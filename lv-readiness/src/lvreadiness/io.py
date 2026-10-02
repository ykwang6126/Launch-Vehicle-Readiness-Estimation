"""Read MATLAB_Input and export the MATLAB v13 workbook and provenance schema."""

import hashlib
import json
from pathlib import Path
import shutil
import subprocess

import pandas as pd
import yaml
from openpyxl import load_workbook

from .validation import InputError


def read_input(path: str | Path) -> pd.DataFrame:
    """Read only MATLAB_Input and return its cell values as a DataFrame."""

    # 1. Confirm the workbook exists before opening it.
    path = Path(path)
    if not path.is_file():
        raise InputError(f"Input workbook not found: {path}")

    # 2. Open Excel in read-only/data-only mode.
    # data_only=True reads stored cell values rather than spreadsheet formulas.
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        # 3. Read only the required MATLAB_Input worksheet.
        if "MATLAB_Input" not in book.sheetnames:
            raise InputError("Required worksheet MATLAB_Input is missing.")
        rows = list(book["MATLAB_Input"].values)
        if not rows:
            raise InputError("MATLAB_Input is empty.")

        # 4. Use the first row as column headers and reject duplicate names.
        headers = list(rows[0])
        if len([h for h in headers if h is not None]) != len(set(h for h in headers if h is not None)):
            raise InputError("MATLAB_Input has duplicate column names.")

        # 5. Return raw worksheet values; validation.py performs model/data checks.
        frame = pd.DataFrame(rows[1:], columns=headers)
        return frame.loc[:, [c is not None for c in frame.columns]]
    finally:
        book.close()


def sha256(path: str | Path) -> str:
    """Fingerprint the exact input file so a run can be traced to its source."""
    with Path(path).open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest() if hasattr(hashlib, 'file_digest') else hashlib.sha256(source.read()).hexdigest()


def git_commit(path: Path) -> str | None:
    """Record the current Git commit in run metadata when Git is available."""

    # This is provenance only; it does not affect the numerical analysis.
    git = shutil.which("git")
    if git is None:
        return None
    try:
        proc = subprocess.run([git, "-C", str(path), "rev-parse", "HEAD"],
                              capture_output=True, text=True, timeout=5, check=False)
        return proc.stdout.strip() if proc.returncode == 0 else None
    except (OSError, subprocess.TimeoutExpired):
        return None


def export_results(run_dir: Path, tables: dict[str, pd.DataFrame], config: dict, metadata: dict) -> dict[str, Path]:
    """Write the analysis tables and reproducibility files for one run."""

    # 1. Export all result tables into the required results.xlsx worksheets.
    workbook = run_dir / "results.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False, na_rep="")

            # Keep text labels as literal strings, even when an ID begins with "=".
            sheet = writer.sheets[name]
            for row_index, values in enumerate(table.itertuples(index=False, name=None), 2):
                for col_index, value in enumerate(values, 1):
                    if isinstance(value, str):
                        sheet.cell(row_index, col_index).data_type = "s"

    # 2. Save the resolved run configuration and provenance metadata separately.
    config_path, metadata_path = run_dir / "run_config.yaml", run_dir / "run_metadata.json"
    config_path.write_text(yaml.safe_dump(config, sort_keys=False), encoding="utf-8")
    metadata_path.write_text(json.dumps(metadata, indent=2, allow_nan=False)+"\n", encoding="utf-8")

    # 3. Return paths so pipeline.py/other Python code can access the outputs.
    return dict(results=workbook, config=config_path, metadata=metadata_path, figures=run_dir/"figures")
