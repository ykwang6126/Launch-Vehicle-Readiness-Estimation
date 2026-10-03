"""Check API/CLI repeatability, exported workbook schemas, figures, and input errors."""

import json
from pathlib import Path
import subprocess
import sys

import numpy as np
import openpyxl
import pandas as pd

from lvreadiness import run_analysis


def test_pipeline_reproducibility_exports_and_cli(synthetic_workbook, tmp_path):
    """Verify identical seeded results, separate run folders, exports, and CLI errors."""
    first = run_analysis(synthetic_workbook, tmp_path / "runs", 3000, 1000, 1)
    second = run_analysis(synthetic_workbook, tmp_path / "runs", 3000, 1000, 1)
    np.testing.assert_array_equal(first.prior_samples, second.prior_samples)
    np.testing.assert_array_equal(first.posterior_samples, second.posterior_samples)
    for key in ("RaterParameters", "PoolSummary", "NodeSummary", "PoolingComparison"):
        pd.testing.assert_frame_equal(
            first.tables[key], second.tables[key], check_exact=True
        )
    assert first.paths["results"] != second.paths["results"]
    book = openpyxl.load_workbook(first.paths["results"], data_only=True)
    assert book.sheetnames == [
        "RunInfo",
        "RaterParameters",
        "PoolSummary",
        "NodeSummary",
        "PoolingComparison",
    ]
    assert [book[s].max_row - 1 for s in book.sheetnames] == [1, 7, 7, 11, 2]
    assert book["NodeSummary"]["J2"].value is None
    assert isinstance(book["NodeSummary"]["D12"].value, float)
    assert len(list(first.paths["figures"].glob("*.png"))) == 12
    metadata = json.loads(first.paths["metadata"].read_text())
    assert (
        metadata["InputMode"] == "linear_pool"
        and metadata["RandomGenerator"] == "PCG64"
    )
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "lvreadiness",
            "run",
            str(synthetic_workbook),
            "--output",
            str(tmp_path / "cli"),
            "--n-prior",
            "3000",
            "--n-posterior",
            "1000",
        ],
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    cli_path = next((tmp_path / "cli").glob("run_*/results.xlsx"))
    pd.testing.assert_frame_equal(
        pd.read_excel(cli_path, sheet_name="NodeSummary"),
        pd.read_excel(first.paths["results"], sheet_name="NodeSummary"),
    )
    failed = subprocess.run(
        [
            sys.executable,
            "-m",
            "lvreadiness",
            "run",
            str(synthetic_workbook),
            "--output",
            str(tmp_path / "bad"),
            "--n-prior",
            "0",
        ],
        capture_output=True,
        text=True,
    )
    assert failed.returncode == 2 and not (tmp_path / "bad").exists()
